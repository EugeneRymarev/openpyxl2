"""Smoke-test the installed wheel using only its required dependencies.

Run with Python's -I option so the checkout cannot shadow the installed package.
"""
import datetime
import io
from pathlib import Path

import openpyxl


checkout = Path(__file__).resolve().parents[2]
assert Path(openpyxl.__file__).resolve().parent != checkout / "openpyxl"
assert not openpyxl.LXML
assert not openpyxl.NUMPY
assert not openpyxl.DEFUSEDXML

values = (42, "Python 3.14", datetime.datetime(2026, 1, 2), "=A1*2")
for write_only in (False, True):
    workbook = openpyxl.Workbook(write_only=write_only)
    sheet = workbook.create_sheet() if write_only else workbook.active
    sheet.append(values)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    for read_only in (False, True):
        stream.seek(0)
        restored = openpyxl.load_workbook(stream, read_only=read_only)
        try:
            assert next(restored.active.values) == values
        finally:
            restored.close()

print("Installed wheel: normal/write-only save and normal/read-only load passed")
