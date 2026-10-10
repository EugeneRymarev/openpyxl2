"""Generate real Excel examples with the imported upstream and current fork.
Run: python doc/readme/build_examples.py
Requires this repository's Python dependencies and the baseline in Git history.
Open the resulting two workbooks in Microsoft Excel to capture the five sheets.
"""

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
BASELINE = "1fa786a9dd5d77754d703f967832f5243e08ba0a"
OPERATIONS = [
    ("insert_rows", 3),
    ("insert_cols", 2),
    ("delete_rows", 2),
    ("delete_cols", 1),
]


def write_examples(target, label):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.views import Selection
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    wb.remove(wb.active)
    results = {}
    line = Side(style="thin", color="FFCBD5E1")

    def view(ws):
        ws.sheet_view.showGridLines = True
        ws.sheet_view.zoomScale = 135
        for column, width in [
            ("A", 4),
            ("B", 22),
            ("C", 16),
            ("D", 18),
            ("E", 22),
            ("F", 18),
            ("G", 4),
        ]:
            ws.column_dimensions[column].width = width
        for row in range(1, 17):
            ws.row_dimensions[row].height = 25

    def footer(ws, operation, expected):
        ws.merge_cells("A12:F12")
        ws["A12"] = f"{label} | {operation}"
        ws["A12"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFFFF")
        ws["A12"].fill = PatternFill(
            "solid", fgColor="FF9C4141" if label == "BEFORE" else "FF217346"
        )
        ws.merge_cells("A14:F14")
        ws["A14"] = expected
        ws["A14"].font = Font(name="Calibri", size=12, color="FF475569")

    for operation, index in OPERATIONS:
        ws = wb.create_sheet(operation)
        view(ws)
        ws.merge_cells("B3:E3")
        ws["B3"] = "OFFICE SUPPLIES"
        ws["B3"].font = Font(name="Calibri", size=18, bold=True, color="FF17365D")
        ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
        for row in [
            (5, ["Item", "Units", "Unit price", "Batch"]),
            (6, ["Pens", 10, 2, "A"]),
            (7, ["Folders", 20, 3, "B"]),
        ]:
            for column, value in enumerate(row[1], 2):
                cell = ws.cell(row[0], column, value)
                cell.font = Font(
                    name="Calibri",
                    size=13,
                    bold=row[0] == 5,
                    color="FFFFFFFF" if row[0] == 5 else "FF17365D",
                )
                cell.fill = PatternFill(
                    "solid", fgColor="FF17365D" if row[0] == 5 else "FFE8EFF7"
                )
                cell.border = Border(left=line, right=line, top=line, bottom=line)
                cell.alignment = Alignment(vertical="center")
        ws["B9"] = "Total units"
        ws["C9"] = "=SUM(C6:C7)"
        for address in ("B9", "C9"):
            ws[address].font = Font(
                name="Calibri", size=16, bold=True, color="FF17365D"
            )
            ws[address].fill = PatternFill("solid", fgColor="FFFFE8A3")
        name = operation + "_units"
        wb.defined_names.add(DefinedName(name, attr_text=f"'{operation}'!$C$6:$C$7"))
        ws.print_area = "B3:E9"
        getattr(ws, operation)(index)
        total = {
            "insert_rows": "C10",
            "insert_cols": "D9",
            "delete_rows": "C8",
            "delete_cols": "B9",
        }[operation]
        results[operation] = {
            "call": f"{operation}({index})",
            "total_cell": total,
            "formula": ws[total].value,
            "merged_ranges": str(ws.merged_cells),
            "defined_name": wb.defined_names[name].attr_text,
            "print_area": str(ws.print_area),
        }
        # Presentation captions are added after the demonstrated operation.
        footer(ws, f"{operation}({index})", "Expected total units: 30 (10 + 20)")
        ws.sheet_view.selection = [Selection(activeCell=total, sqref=total)]

    ws = wb.create_sheet("merge_cells")
    view(ws)
    ws["B1"] = "STYLES ASSIGNED AFTER MERGING"
    ws["B1"].font = Font(name="Calibri", size=17, bold=True, color="FF17365D")
    side = Side(style="medium", color="FF2563EB")
    for area, anchor, value in [
        ("B3:E6", "B3", "Merged B3:E6"),
        ("B9:E9", "B9", "Merged B9:E9"),
    ]:
        ws.merge_cells(area)
        cell = ws[anchor]
        cell.value = value
        cell.font = Font(name="Calibri", size=18, bold=True, color="FF17365D")
        cell.fill = PatternFill("solid", fgColor="FFDBEAFE")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
    results["merge_cells"] = {
        "C4_fill": ws["C4"].fill.fgColor.rgb,
        "E6_right": getattr(ws["E6"].border.right, "style", None),
        "E6_bottom": getattr(ws["E6"].border.bottom, "style", None),
    }
    footer(
        ws,
        "merge_cells + anchor styles",
        "Expected: a complete blue outline around each merged range",
    )
    ws.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    wb.active = 0
    wb.save(target)
    target.with_suffix(".json").write_text(
        json.dumps(results, indent=2), encoding="utf-8"
    )
    print(target)


def main():
    output = HERE / "workbooks"
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="openpyxl-readme-") as temporary:
        baseline = Path(temporary) / "upstream"
        baseline.mkdir()
        archive = Path(temporary) / "upstream.zip"
        subprocess.run(
            ["git", "archive", "--format=zip", "-o", str(archive), BASELINE],
            cwd=ROOT,
            check=True,
        )
        with zipfile.ZipFile(archive) as source:
            source.extractall(baseline)
        for source, label, filename in [
            (baseline, "BEFORE", "before.xlsx"),
            (ROOT, "AFTER", "after.xlsx"),
        ]:
            env = os.environ.copy()
            env["PYTHONPATH"] = str(source)
            subprocess.run(
                [
                    sys.executable,
                    str(Path(__file__).resolve()),
                    "--worker",
                    str(output / filename),
                    label,
                ],
                cwd=source,
                env=env,
                check=True,
            )
    current = json.loads((output / "after.json").read_text())
    old = json.loads((output / "before.json").read_text())
    expected = {
        "insert_rows": "=SUM(C7:C8)",
        "insert_cols": "=SUM(D6:D7)",
        "delete_rows": "=SUM(C5:C6)",
        "delete_cols": "=SUM(B6:B7)",
    }
    for name, formula in expected.items():
        assert current[name]["formula"] == formula, (name, current[name])
        assert old[name]["formula"] == "=SUM(C6:C7)", (name, old[name])
    assert current["merge_cells"]["E6_bottom"] == "medium"
    assert old["merge_cells"]["E6_bottom"] is None
    (HERE / "provenance.json").write_text(
        json.dumps(
            {
                "upstream": BASELINE,
                "fork": subprocess.check_output(
                    ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
                ).strip(),
                "before": old,
                "after": current,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print("All five old/new examples verified.")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--worker":
        write_examples(Path(sys.argv[2]), sys.argv[3])
    else:
        main()
